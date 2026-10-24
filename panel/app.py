"""Panel de control local: edita config/control_plataforma.xlsx y ejecuta ops.py desde el navegador.
Uso: python panel/app.py   (o python ops.py panel)   -> http://127.0.0.1:8765
Solo escucha en 127.0.0.1: no lo expongas a la red. No sustituye revisar `python ops.py plan` antes de aplicar.
"""
import datetime as dt
import json
import subprocess
import sys
import threading
import time
import uuid
from pathlib import Path

from fastapi import FastAPI
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "config/control_plataforma.xlsx"

app = FastAPI(title="Panel STYLUS")
JOBS: dict[str, dict] = {}  # id -> {proc, lines: [], done: bool, comando}

# Comandos del orquestador expuestos en el panel: (clave, etiqueta, flags, peligroso)
COMANDOS = [
    ("config", "Validar / generar configuración", [], False),
    ("estado", "Ver estado", [], False),
    ("bootstrap", "1. Preparar proyecto (APIs, red, estado)", [], False),
    ("plan", "2. Revisar plan de Terraform", [], False),
    ("apply", "3. Crear infraestructura", ["--yes"], True),
    ("subir-datos", "4. Subir datos H&M a GCS", [], False),
    ("bronze", "5. Cargar Bronze", [], False),
    ("silver", "6. Procesar Silver (Spark)", [], False),
    ("gold", "7. Calcular Gold y vistas BI", [], False),
    ("api", "8. Desplegar API + catálogo", [], False),
    ("verificar", "9. Verificar BigQuery vs Excel", [], False),
    ("github", "Configurar variable de GitHub Actions", [], False),
    ("costo-real", "Ver costo real por pedido", [], False),
    ("destroy", "Destruir infraestructura", ["--yes"], True),
]


EDITABLE_FILL = "FFF2CC"  # mismo amarillo claro usado al construir config/control_plataforma.xlsx


def _cell_style(ws, row):
    fill = ws.cell(row=row, column=2).fill
    rgb = fill.fgColor.rgb if fill and fill.fgColor else None
    return isinstance(rgb, str) and rgb.upper().endswith(EDITABLE_FILL)


def leer_parametros():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Parametros"]
    head = [c.value for c in ws[1]]
    i_clave, i_valor, i_tipo, i_desc = (head.index(h) for h in ("clave", "valor", "tipo", "descripcion"))
    out, seccion = [], "General"
    for row in range(2, ws.max_row + 1):
        vals = [ws.cell(row=row, column=c + 1).value for c in range(len(head))]
        clave = vals[i_clave]
        if clave is None:
            continue
        if isinstance(clave, str) and clave.startswith("#"):
            seccion = clave.lstrip("# ").strip()
            continue
        out.append({
            "clave": clave, "valor": vals[i_valor], "tipo": vals[i_tipo] or "texto",
            "descripcion": vals[i_desc] or "", "seccion": seccion,
            "editable": _cell_style(ws, row), "fila": row,
        })
    wb.close()
    return out


def escribir_parametros(cambios: dict):
    """cambios: {clave: valor_nuevo}. Solo toca celdas de valor marcadas editable; conserva fórmulas y formato."""
    from openpyxl import load_workbook
    params = {p["clave"]: p for p in leer_parametros()}
    rechazados = [k for k in cambios if k not in params or not params[k]["editable"]]
    wb = load_workbook(XLSX)
    ws = wb["Parametros"]
    for clave, valor in cambios.items():
        if clave in rechazados:
            continue
        ws.cell(row=params[clave]["fila"], column=2).value = valor
    wb.save(XLSX)
    wb.close()
    return rechazados


class Cambios(BaseModel):
    valores: dict


@app.get("/", response_class=HTMLResponse)
def index():
    return (Path(__file__).parent / "static" / "index.html").read_text(encoding="utf-8")


@app.get("/api/comandos")
def comandos():
    return [{"comando": c, "etiqueta": e, "peligroso": d} for c, e, _, d in COMANDOS]


@app.get("/api/parametros")
def parametros():
    try:
        return {"ok": True, "items": leer_parametros()}
    except PermissionError:
        return {"ok": False, "error": "No se pudo abrir el Excel: cierra config/control_plataforma.xlsx en Excel e inténtalo de nuevo."}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/api/parametros")
def guardar(cambios: Cambios):
    try:
        rechazados = escribir_parametros(cambios.valores)
        return {"ok": True, "rechazados": rechazados}
    except PermissionError:
        return {"ok": False, "error": "El Excel está abierto en otro programa (Excel). Cierralo e inténtalo de nuevo."}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def _run_job(job_id, comando, flags):
    proc = subprocess.Popen(
        [sys.executable, str(ROOT / "ops.py"), comando, *flags],
        cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, bufsize=1,
    )
    JOBS[job_id]["proc"] = proc
    for line in proc.stdout:
        JOBS[job_id]["lines"].append(line.rstrip())
    proc.wait()
    JOBS[job_id]["done"] = True
    JOBS[job_id]["codigo"] = proc.returncode


@app.post("/api/ejecutar/{comando}")
def ejecutar(comando: str):
    cfg = {c: (c, e, f, d) for c, e, f, d in COMANDOS}.get(comando)
    if not cfg:
        return {"ok": False, "error": "comando desconocido"}
    _, _, flags, _ = cfg
    job_id = uuid.uuid4().hex[:12]
    JOBS[job_id] = {"lines": [], "done": False, "comando": comando, "inicio": dt.datetime.now().isoformat()}
    threading.Thread(target=_run_job, args=(job_id, comando, flags), daemon=True).start()
    return {"ok": True, "job_id": job_id}


@app.get("/api/log/{job_id}")
def log_stream(job_id: str):
    def gen():
        enviado = 0
        while True:
            job = JOBS.get(job_id)
            if job is None:
                yield "data: " + json.dumps({"error": "job no encontrado"}) + "\n\n"
                return
            nuevas = job["lines"][enviado:]
            enviado += len(nuevas)
            for linea in nuevas:
                yield "data: " + json.dumps({"linea": linea}) + "\n\n"
            if job["done"] and enviado >= len(job["lines"]):
                yield "data: " + json.dumps({"fin": True, "codigo": job.get("codigo", 0)}) + "\n\n"
                return
            time.sleep(0.5)
    return StreamingResponse(gen(), media_type="text/event-stream")


if __name__ == "__main__":
    import uvicorn
    print("Panel de control: http://127.0.0.1:8765  (Ctrl+C para detener)")
    uvicorn.run(app, host="127.0.0.1", port=8765)
