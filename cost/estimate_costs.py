"""Calculadora de costos GCP automatizada.

  python cost/estimate_costs.py                 # genera cost/output/cost_estimate.xlsx + resumen.md
  python cost/estimate_costs.py --refresh-prices # actualiza precios desde Cloud Billing Catalog API
                                                 # (requiere `gcloud auth login`); reescribe assumptions.yaml
Toda la lógica del Excel queda en fórmulas: cambiar un precio o un volumen recalcula todo.
"""
import argparse
import datetime as dt
import json
import re
import subprocess
import urllib.request
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
CFG = HERE / "assumptions.yaml"
OUT = HERE / "output"
API = "https://cloudbilling.googleapis.com/v1"


def compute(cfg):
    """Mismo cálculo que las fórmulas del Excel (para documentos y CI)."""
    p = cfg["precios"]
    res = {}
    for esc, e in cfg["escenarios"].items():
        det = {k: max(0.0, e["uso"].get(k, 0) - v["gratis"]) * v["precio"] for k, v in p.items()}
        total = sum(det.values())
        res[esc] = {"total": total, "pedidos": e["pedidos_mes"], "por_pedido": total / e["pedidos_mes"], "detalle": det}
    for bloque in ("anio1", "curso"):
        res[bloque] = sum(res[k]["total"] * m for k, m in cfg[bloque].items())
    return res


# ---------- refresco de precios ----------
def _get(url, token):
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def _paged(url, token, key):
    tok = ""
    while True:
        d = _get(f"{url}{'&' if '?' in url else '?'}pageSize=5000&pageToken={tok}", token)
        yield from d.get(key, [])
        tok = d.get("nextPageToken")
        if not tok:
            return


def refresh_prices(cfg):
    token = subprocess.run(["gcloud", "auth", "print-access-token"], capture_output=True, text=True, check=True).stdout.strip()
    services = {s["displayName"]: s["name"] for s in _paged(f"{API}/services", token, "services")}
    sku_cache = {}
    for key, item in cfg["precios"].items():
        name = next((n for disp, n in services.items() if disp.lower().startswith(item["servicio"].split()[0].lower() + " ")
                     or disp.lower() == item["servicio"].lower()), None)
        if not name:
            print(f"[sin servicio] {key}")
            continue
        if name not in sku_cache:
            sku_cache[name] = list(_paged(f"{API}/{name}/skus", token, "skus"))
        rx = re.compile(item["sku_regex"], re.I)
        hit = next((s for s in sku_cache[name] if rx.search(s["description"])
                    and (cfg["region"] in s.get("serviceRegions", []) or "global" in s.get("serviceRegions", []))), None)
        if not hit:
            print(f"[sin SKU] {key} ({item['sku_regex']}) -> se mantiene {item['precio']}")
            continue
        rate = hit["pricingInfo"][0]["pricingExpression"]["tieredRates"][-1]["unitPrice"]
        new = int(rate.get("units", 0)) + rate.get("nanos", 0) / 1e9
        print(f"[ok] {key}: {item['precio']} -> {new}  ({hit['description']})")
        item["precio"] = round(new, 9)
    cfg["fecha_precios"] = dt.date.today().isoformat()
    CFG.write_text(yaml.safe_dump(cfg, allow_unicode=True, sort_keys=False, width=200), encoding="utf-8")


# ---------- Excel con fórmulas ----------
BLUE = Font(name="Arial", color="0000FF")
BOLD = Font(name="Arial", bold=True)
HDR = Font(name="Arial", bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BASE = Font(name="Arial")
USD = '"$"#,##0.00;("$"#,##0.00);"-"'


def header(ws, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HDR, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"


def build_xlsx(cfg, path):
    p, esc = cfg["precios"], cfg["escenarios"]
    keys, names = list(p), list(esc)
    n = len(keys)
    wb = Workbook()

    ws = wb.active
    ws.title = "Precios"
    header(ws, ["Clave", "Servicio", "Unidad", "Precio USD", "Gratis/mes", "Fuente"])
    for r, k in enumerate(keys, 2):
        v = p[k]
        ws.append([k, v["servicio"], v["unidad"], v["precio"], v["gratis"], v["fuente"]])
        for c in (4, 5):
            ws.cell(r, c).font, ws.cell(r, c).fill = BLUE, INPUT_FILL
    ws.cell(n + 3, 1, f"Precios lista {cfg['region']}, consultados {cfg['fecha_precios']}. Celdas amarillas = editables. "
                      "Estimación, no cotización: validar en https://cloud.google.com/products/calculator").font = BASE

    wu = wb.create_sheet("Uso")
    header(wu, ["Clave", "Unidad"] + names)
    for r, k in enumerate(keys, 2):
        wu.append([k, p[k]["unidad"]] + [esc[e]["uso"].get(k, 0) for e in names])
        for c in range(3, 3 + len(names)):
            wu.cell(r, c).font, wu.cell(r, c).fill = BLUE, INPUT_FILL
    rp = n + 3
    wu.cell(rp, 1, "pedidos_mes").font = BOLD
    for j, e in enumerate(names, 3):
        c = wu.cell(rp, j, esc[e]["pedidos_mes"])
        c.font, c.fill = BLUE, INPUT_FILL
    wu.cell(rp + 1, 1, "descripcion").font = BOLD
    for j, e in enumerate(names, 3):
        wu.cell(rp + 1, j, esc[e]["descripcion"]).alignment = Alignment(wrap_text=True)

    wc = wb.create_sheet("Costo mensual")
    header(wc, ["Clave", "Servicio"] + names)
    for r, k in enumerate(keys, 2):
        wc.cell(r, 1, k)
        wc.cell(r, 2, f"=Precios!B{r}")
        for j in range(len(names)):
            col = get_column_letter(3 + j)
            cell = wc.cell(r, 3 + j, f"=MAX(0,Uso!{col}{r}-Precios!$E{r})*Precios!$D{r}")
            cell.number_format = USD
    tr = n + 2
    wc.cell(tr, 1, "TOTAL MENSUAL").font = BOLD
    wc.cell(tr + 1, 1, "Pedidos/mes").font = BOLD
    wc.cell(tr + 2, 1, "Costo por pedido").font = BOLD
    for j in range(len(names)):
        col = get_column_letter(3 + j)
        wc.cell(tr, 3 + j, f"=SUM({col}2:{col}{n + 1})").number_format = USD
        wc.cell(tr + 1, 3 + j, f"=Uso!{col}{rp}").number_format = "#,##0"
        wc.cell(tr + 2, 3 + j, f"=IF({col}{tr + 1}>0,{col}{tr}/{col}{tr + 1},0)").number_format = '"$"0.00000'
        for rr in (tr, tr + 2):
            wc.cell(rr, 3 + j).font = BOLD

    wr = wb.create_sheet("Resumen", 0)
    header(wr, ["Concepto", "Meses", "USD"])
    wr.append(["Escenario", "", "Total mensual"])
    for j, e in enumerate(names):
        col = get_column_letter(3 + j)
        wr.append([e, "", f"='Costo mensual'!{col}{tr}"])
    r0 = wr.max_row + 2
    blocks = [("Año 1 (producción)", cfg["anio1"]), ("Costo del curso", cfg["curso"])]
    for i, (label, mix) in enumerate(blocks):
        terms = "+".join(f"'Costo mensual'!{get_column_letter(3 + names.index(k))}{tr}*{m}" for k, m in mix.items())
        wr.cell(r0 + i, 1, label).font = BOLD
        wr.cell(r0 + i, 2, " + ".join(f"{m}×{k}" for k, m in mix.items()))
        wr.cell(r0 + i, 3, f"={terms}").font = BOLD
    pr = names.index("produccion")
    wr.cell(r0 + 3, 1, "Costo por pedido (producción)").font = BOLD
    wr.cell(r0 + 3, 3, f"='Costo mensual'!{get_column_letter(3 + pr)}{tr + 2}").number_format = '"$"0.00000'
    wr.cell(r0 + 5, 1, "Año 1 por servicio").font = BOLD
    servicios = sorted({v["servicio"] for v in p.values()})
    for i, s in enumerate(servicios):
        terms = "+".join(f"SUMIF('Costo mensual'!$B$2:$B${n + 1},A{r0 + 6 + i},'Costo mensual'!{get_column_letter(3 + names.index(k))}$2:{get_column_letter(3 + names.index(k))}${n + 1})*{m}"
                         for k, m in cfg["anio1"].items())
        wr.cell(r0 + 6 + i, 1, s)
        wr.cell(r0 + 6 + i, 3, f"={terms}")
    for row in wr.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            if c.number_format == "General":
                c.number_format = USD

    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.font == Font():
                    c.font = BASE
        for i, w in enumerate([28, 30, 22, 22, 22, 22, 60][: sh.max_column], 1):
            sh.column_dimensions[get_column_letter(i)].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh-prices", action="store_true")
    a = ap.parse_args()
    cfg = yaml.safe_load(CFG.read_text(encoding="utf-8"))
    if a.refresh_prices:
        refresh_prices(cfg)
    res = compute(cfg)
    build_xlsx(cfg, OUT / "cost_estimate.xlsx")
    lines = [f"# Estimación de costos GCP ({cfg['region']}, precios {cfg['fecha_precios']})", "",
             "| Escenario | USD/mes | Pedidos/mes | USD/pedido |", "|---|---:|---:|---:|"]
    for k, v in res.items():
        if isinstance(v, dict):
            lines.append(f"| {k} | {v['total']:,.2f} | {v['pedidos']:,} | {v['por_pedido']:.5f} |")
    lines += ["", f"**Año 1 producción:** USD {res['anio1']:,.2f}  ", f"**Costo del curso (dev+demo):** USD {res['curso']:,.2f}"]
    (OUT / "resumen.md").write_text("\n".join(lines), encoding="utf-8")
    (OUT / "resumen.json").write_text(json.dumps(res, indent=2), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
