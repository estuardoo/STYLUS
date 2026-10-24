"""Genera docs/gestion/Matriz_Auditoria_Proyecto.xlsx y docs/DEPLOY.md desde una sola fuente.
Uso: python docs/_generadores/build_matriz.py   (desde la raíz del repo)
Ojo: sobrescribe los estados editados a mano en la Matriz de auditoría; actualizar aquí primero."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
E, W, N, G, T = "Estuardo Oliver", "Winston Quispe", "Nicole Arenas", "Gianmarco Mejia", "Todo el equipo"
F1, F2, F3 = "24-oct-2026", "07-nov-2026", "20-nov-2026"
H, L, P = "Hecho", "Listo para desplegar", "Pendiente"

PERFILES = [
    (E, "ML / Analytics", "Chatbot con Vertex AI Gemini, analítica, recomendaciones; integra documentos y demo."),
    (W, "Cloud Computing", "Todo lo nube: Terraform, IAM, GCS, Pub/Sub, Dataflow (ingesta streaming), CI/CD, costos, DRP."),
    (N, "Datos + API + Modelado", "Bronze/Silver/Gold, calidad, modelo estrella, FastAPI y base transaccional."),
    (G, "Dashboard / BI", "KPIs, Looker Studio, calculadora y tablero de costos, diagramas y slides."),
]

# ID, paso, fase, nombre, acción, entregable, comando, depende, responsable, apoyo, estado, evidencia
TAREAS = [
    ("P1-01", 1, "Datos", "Auditoría de datos H&M", "Perfilar filas, nulos, fechas, cardinalidad y cobertura", "Word §4", "", "", N, E, H, "31.8M líneas, 9.08M pedidos, 1.37M clientes, 105,542 artículos"),
    ("P1-02", 1, "Gestión", "Ordenar carpeta y repositorio", "Estructura por capas; Amazon/zip archivados; secretos fuera", "README.md", "", "", E, W, H, "data/_archive y .secrets/ fuera de Git"),
    ("P1-03", 1, "Seguridad", "Proteger datos y secretos", ".gitignore + regla CI que bloquea CSV/ZIP/tfstate/credenciales", ".gitignore, .github/workflows/ci.yml", "", "", W, "", H, "Regla 'Sin datos ni secretos versionados'"),
    ("P1-04", 1, "Diseño", "Justificar stack GCP", "Decisión por capa con alternativa descartada", "Word §5", "", "", W, E, H, "Tabla de ADR"),
    ("P1-05", 1, "Diseño", "Diagrama end-to-end", "Fuentes, ingesta, almacenamiento, cómputo, exposición", "docs/architecture/arquitectura.png", "dot -Tpng docs/architecture/arquitectura_vertical.dot -o docs/architecture/arquitectura.png", "", G, W, H, "Graphviz regenerable"),
    ("P1-06", 1, "Diseño", "Integración multifuente", "Batch H&M + replay 2026; doble reloj source_t_dat/event_time", "Word §6", "", "P1-01", N, W, H, ""),
    ("P1-07", 1, "Diseño", "Ingesta y almacenamiento", "GCS medallion + BigQuery + Cloud SQL; formatos y particiones", "Word §7", "", "P1-06", N, W, H, ""),
    ("P1-08", 1, "Diseño", "Escalabilidad Black Friday", "Autoescalado, desacople, degradación controlada", "Word §8", "", "", W, N, H, "Escenario 5x en calculadora"),
    ("P1-09", 1, "Costos", "Calculadora y costo por pedido", "YAML -> Excel con fórmulas; 4 escenarios; año 1", "docs/entrega1/Calculadora_Costos_GCP.xlsx", "python cost/estimate_costs.py", "", G, W, H, "Año 1 USD 5,482; USD 0.00098/pedido (Gemini 3.1 Flash-Lite)"),
    ("P1-10", 1, "Costos", "Automatizar control de costos", "Precios vía API, budget con alertas, costo real por pedido", "cost/, infra/terraform/main.tf", "python cost/estimate_costs.py --refresh-prices", "P2-01", W, G, L, "Probado local; activar billing export en consola"),
    ("P1-11", 1, "Resiliencia", "DRP / HA", "RPO/RTO por componente y runbook", "Word §10, docs/runbooks/drp.md", "", "", W, N, H, ""),
    ("P1-12", 1, "BI", "Catálogo de KPIs", "Metas de datos, streaming, modelo, API, negocio, costo", "Word §11", "", "", G, E, H, ""),
    ("P1-13", 1, "Datos", "Ingesta Bronze con manifiesto", "SHA-256 y filas por archivo; rsync idempotente", "ingestion/batch/upload_raw.py", "python ingestion/batch/upload_raw.py --manifest-only", "", N, W, H, "Manifiesto real en data/manifests/"),
    ("P1-14", 1, "Infra", "Terraform base", "Buckets, BigQuery, Pub/Sub+DLQ, IAM, Artifact Registry, budget (valores desde el JSON de control)", "infra/terraform/", "python ops.py plan", "P1-20", W, "", L, "validate OK; tipos de variables probados con el JSON generado"),
    ("P1-15", 1, "Gestión", "Documento Word Entrega 1", "Examen + P1 profesional", "docs/entrega1/*.docx (+PDF)", "", "P1-01..P1-14", E, T, H, ""),
    ("P1-16", 1, "Gestión", "Matriz de auditoría y roles", "Tareas, pasos, despliegue y variables", "docs/gestion/Matriz_Auditoria_Proyecto.xlsx", "python docs/_generadores/build_matriz.py", "", E, T, H, "Este archivo"),
    ("P1-17", 1, "Infra", "Crear proyecto GCP", "Proyecto, facturación/créditos, accesos IAM del equipo", "Proyecto GCP", "ver hoja Despliegue D01-D04", "", W, "", P, ""),
    ("P1-18", 1, "Gestión", "Publicar repositorio GitHub", "Crear repo privado, push y CI en verde", "URL del repo", "git push -u origin main", "P1-17", W, T, P, "Git inicializado localmente"),
    ("P1-20", 1, "Infra", "JSON de control de la plataforma", "Parámetros, APIs, cuentas, permisos, usuarios, datasets, tablas, columnas y secretos en un solo JSON validado, editable con el panel web", "config/control_plataforma.json", "python ops.py config", "", W, T, H, "30 tablas, 240 columnas, 28 permisos, 6 usuarios app; CI verifica generados"),
    ("P1-21", 1, "Infra", "Orquestador de despliegue", "Un comando por paso, multiplataforma, con log matricial de tiempos", "ops.py", "python ops.py todo", "P1-20", W, N, L, "config/estado/local probados; comandos de nube pendientes de GCP"),
    ("P1-19", 1, "Gestión", "Revisión y envío Entrega 1", "Revisión cruzada del Word y envío antes de 23:59", "Word enviado", "", "P1-15", E, T, P, ""),

    ("P2-01", 2, "Infra", "Desplegar infraestructura dev", "Bootstrap de estado, plan y apply", "Recursos en GCP", "python ops.py bootstrap; python ops.py apply", "P1-17", W, "", P, ""),
    ("P2-02", 2, "Datos", "Cargar Bronze (GCS + BigQuery)", "CSV primero, imágenes después; tablas externas y nativa particionada", "retail_bronze.*", "python ops.py subir-datos; python ops.py bronze", "P2-01", N, W, P, "SQL y script listos"),
    ("P2-03", 2, "Datos", "Silver con Spark", "Tipado, nulos, cuarentena con motivo, Parquet particionado, reporte conciliado", "pipelines/batch/silver_hm.py", "python ops.py silver", "P2-02", N, W, L, "Local: 546,741 filas, 3 rechazos inyectados detectados, conciliado; corregido fn/active"),
    ("P2-04", 2, "Datos", "Silver en BigQuery + vista de eventos", "LOAD DATA de Parquet; eventos deduplicados por event_id", "sql/silver/01_load_silver.sql", "python ops.py silver", "P2-03", N, "", L, ""),
    ("P2-05", 2, "Modelado", "Gold: modelo estrella", "fact_sales, fact_orders, dim_article/customer/date/channel", "sql/gold/01_gold_model.sql", "python ops.py gold", "P2-04", N, E, L, ""),
    ("P2-06", 2, "Streaming", "Replay de eventos 2026", "Transacciones reales re-emitidas con event_id idempotente", "ingestion/streaming/replay_publisher.py", "python ops.py replay", "P2-01", W, "", L, "Dry-run sobre datos reales OK"),
    ("P2-07", 2, "Streaming", "Ruta económica Pub/Sub -> BigQuery", "Suscripción BigQuery con DLQ", "infra/terraform/main.tf", "", "P2-01", W, "", L, ""),
    ("P2-08", 2, "Streaming", "Pipeline Dataflow", "Validación, dedup, ventanas 1 min, datos tardíos, DLQ, lag p50/p95", "pipelines/streaming/dataflow_events.py", "python ops.py dataflow-iniciar / dataflow-detener", "P2-06", W, N, L, "DirectRunner: 350 mensajes -> 300 únicos, 2 a DLQ"),
    ("P2-09", 2, "ML", "Popularidad y catálogo de servicio", "reco_popular y v_catalog_serving en Gold", "sql/gold/01_gold_model.sql", "python ops.py gold", "P2-05", E, N, L, ""),
    ("P2-10", 2, "ML", "Recomendador baseline + evaluación", "Split temporal, MAP@12: popularidad vs recompra", "docs/evaluacion_recomendador.md", "python pipelines/reco/baseline_eval.py", "", E, "", H, "MAP@12 0.00875 vs 0.02622 (+200%)"),
    ("P2-11", 2, "Datos", "Diccionario de datos", "Columnas, reglas y linaje Silver/Gold/ops/OLTP", "docs/diccionario_datos.md", "", "", N, "", H, ""),
    ("P2-12", 2, "BI", "Vistas para tableros", "Ventas mensuales, top artículos, recurrencia, salud streaming", "sql/gold/02_kpi_views.sql", "python ops.py gold", "P2-05", G, N, L, ""),
    ("P2-13", 2, "BI", "Dashboard borrador", "Primer tablero Looker Studio sobre las vistas Gold", "Enlace Looker Studio", "", "P2-12", G, "", P, ""),
    ("P2-14", 2, "Gestión", "Informe Entrega 2", "Evidencias: conteos, reporte de calidad, capturas", "Word Entrega 2", "", "P2-02..P2-13", E, T, P, ""),

    ("P3-01", 3, "API", "API FastAPI", "Catálogo, recomendaciones, asistente, carrito, checkout, wallet, pedidos, métricas, OpenAPI /docs", "api/main.py", "python ops.py local", "", N, W, L, "Pruebas automáticas OK"),
    ("P3-02", 3, "API", "Wallet transaccional", "Ledger append-only, ACID, idempotencia, sin sobregiro, confirmación explícita", "api/schema.sql", "python -m unittest discover -s tests", "", N, "", L, "Test: doble clic = 1 cobro; 409 sin saldo"),
    ("P3-03", 3, "ML", "Asistente Vertex AI Gemini", "Gemini 3.1 Flash-Lite interpreta; backend busca en catálogo real; respaldo por reglas", "api/assistant.py", "config/control_plataforma.json: assistant_mode = gemini", "P2-01", E, N, L, "Reglas probadas con catálogo real; Gemini requiere GCP"),
    ("P3-04", 3, "ML", "Recomendaciones en API", "Perfil por grupo + popularidad; cold-start explícito", "api/main.py", "", "", E, "", L, ""),
    ("P3-05", 3, "API", "Web tienda", "Chat, productos con imagen, carrito, confirmación, saldo, pedidos", "api/static/index.html", "python ops.py local", "", N, G, L, ""),
    ("P3-06", 3, "Infra", "Cloud SQL + Cloud Run + secretos", "Terraform de base, servicio, secretos y acceso público", "infra/terraform/cloudsql.tf, cloudrun.tf", "python ops.py apply", "P2-01", W, "", L, "tofu validate OK"),
    ("P3-07", 3, "Infra", "CI/CD sin llaves", "Workload Identity Federation + workflow de despliegue", "infra/terraform/cicd.tf, .github/workflows/deploy.yml", "python ops.py github; Actions > deploy-api", "P1-18", W, "", L, ""),
    ("P3-08", 3, "Datos", "Catálogo Gold -> Cloud SQL", "Cargar v_catalog_serving con Cloud SQL Auth Proxy", "Tabla articles en Cloud SQL", "python ops.py catalogo", "P2-09, P3-06", N, W, P, "Local: 105,542 artículos cargados"),
    ("P3-09", 3, "BI", "Dashboard gerencial", "Ventas, ticket medio, top, recurrencia, recomendador", "Looker Studio", "", "P2-12", G, E, P, ""),
    ("P3-10", 3, "BI", "Dashboard operativo y de costos", "Eventos/min, lag, DLQ, costo real y USD/pedido", "Looker Studio", "", "P2-08, P1-10", G, W, P, ""),
    ("P3-11", 3, "Resiliencia", "Prueba de recuperación", "Clonar Cloud SQL por PITR y restaurar objeto GCS; medir RTO", "docs/runbooks/drp.md", "", "P3-06", W, N, P, "Runbook listo"),
    ("P3-12", 3, "Costos", "Costo real vs estimado", "Billing export, vista de costo real, recalibrar supuestos", "cost/real_cost_per_order.sql", "", "P1-10", G, W, P, ""),
    ("P3-13", 3, "Gestión", "Guion de demo end-to-end", "Caso bebé: asistente -> carrito -> compra -> evento -> Gold -> tablero", "docs/demo_script.md", "", "", E, T, H, ""),
    ("P3-14", 3, "Gestión", "Slides (8 puntos del PDF)", "Resumen, situación, KPI, productos, supuestos, arquitectura, calculadora, demo, conclusiones", "docs/slides/", "", "P3-01..P3-12", G, T, P, ""),
    ("P3-15", 3, "Gestión", "Documento final y trazabilidad", "Word final con evidencias y matriz requisito -> evidencia", "docs/", "", "P3-14", E, T, P, ""),
    ("P3-16", 3, "Costos", "Apagado post-demo", "Detener Dataflow; destruir recursos demo", "Checklist", "python ops.py dataflow-detener; python ops.py destroy", "P3-15", W, "", P, ""),
]

# Paso, #, acción, comando, responsable, insumo del JSON de control, verificación
DESPLIEGUE = [
    (1, "D01", "Instalar herramientas", "Google Cloud CLI, Terraform >= 1.6, Python 3.10-3.12, git (o usar Cloud Shell, que ya las trae)", W, "", "gcloud --version; terraform -version"),
    (1, "D02", "Crear proyecto y vincular facturación", "gcloud projects create <project_id>; gcloud billing projects link <project_id> --billing-account=<cuenta>", W, "Parametros: project_id, billing_account", "Proyecto con facturación en consola"),
    (1, "D03", "Completar el JSON de control", "Editar con `python ops.py panel` (o a mano) los campos marcados editables de config/control_plataforma.json: mínimo project_id y emails del equipo", T, "Parametros, Usuarios", "python ops.py config sin errores"),
    (1, "D04", "Autenticarse", "gcloud auth login; gcloud auth application-default login (en Cloud Shell no hace falta)", T, "", "gcloud config list"),
    (1, "D05", "Instalar dependencias", "pip install -r requirements.txt", T, "", "python ops.py estado"),
    (1, "D06", "Publicar repo", "git remote add origin https://github.com/OWNER/REPO.git; git push -u origin main", W, "Parametros: github_repo", "Workflow ci en verde"),
    (2, "D07", "Preparar proyecto", "python ops.py bootstrap", W, "APIs, Parametros: region, dataproc_subnet, bucket_tfstate_suffix", "APIs activas, bucket de estado, terraform init"),
    (2, "D08", "Revisar plan", "python ops.py plan", W, "Todas las hojas", "Plan sin errores; recursos esperados"),
    (2, "D09", "Crear infraestructura", "python ops.py apply", W, "Todas las hojas", "Buckets, datasets, tablas, Pub/Sub, IAM, Cloud SQL, Cloud Run"),
    (2, "D10", "Subir datos H&M (desde la PC con data/raw)", "python ops.py subir-datos  (imágenes: --con-imagenes)", N, "Parametros: local_raw_dir, bucket_raw_suffix", "gs://<bucket raw>/hm/*.csv"),
    (2, "D11", "Bronze", "python ops.py bronze", N, "Datasets, Tablas", "31,788,324 filas en bronze.transactions"),
    (2, "D12", "Silver (Serverless Spark)", "python ops.py silver", N, "Parametros: dataproc_runtime", "_quality_report conciliado=true"),
    (2, "D13", "Gold y vistas BI", "python ops.py gold", N, "Parametros: price_demo_factor", "Tablas fact_/dim_ y vistas v_*"),
    (2, "D14", "Streaming de prueba", "python ops.py dataflow-iniciar; python ops.py replay; python ops.py dataflow-detener", W, "Parametros: dataflow_*, replay_*", "Filas en ops.stream_metrics_1m y silver.v_events"),
    (3, "D15", "API, web y catálogo", "python ops.py api", N, "Parametros: cloudrun_*, assistant_mode, gemini_model; Usuarios (app)", "URL pública; /ready con 105,542 artículos"),
    (3, "D16", "Verificar contra el JSON de control", "python ops.py verificar", T, "Tablas, Columnas", "Matriz tabla por tabla: OK / FALTA / DIFERENTE"),
    (3, "D17", "Despliegue desde GitHub", "python ops.py github -> crear variable GCP_WIF_PROVIDER -> Actions > deploy-api", W, "Parametros: github_repo", "Workflow deploy-api en verde"),
    (3, "D18", "Tableros", "Looker Studio > Agregar datos > BigQuery > gold.v_* y ops.*", G, "Tablas (vistas)", "Tablero gerencial y operativo"),
    (3, "D19", "Costos reales", "Consola: exportar facturación a BigQuery (dataset ops) -> parámetro billing_export_table -> python ops.py costo-real", W, "Parametros: billing_export_table", "ops.v_costo_real_mensual con datos"),
    (3, "D20", "Apagado", "python ops.py dataflow-detener; python ops.py destroy", W, "Parametros: force_destroy_buckets", "Facturación diaria ~0"),
]

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HDR, FILL = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
BRD = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
PASO_FILL = {1: "DDEBF7", 2: "E2EFDA", 3: "FCE4D6"}


def sheet(wb, title, cols, widths, rows, index=None):
    ws = wb.create_sheet(title, index) if wb.worksheets and wb.active.title != "Sheet" else wb.active
    ws.title = title
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(1, i, c)
        cell.font, cell.fill, cell.border = HDR, FILL, BRD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border, cell.alignment = FONT, BRD, WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    return ws


def build_xlsx(out):
    wb = Workbook()
    rol = {p[0]: p[1] for p in PERFILES}
    orden = {}
    rows = []
    for t in TAREAS:
        tid, paso, fase, nom, acc, ent, cmd, dep, resp, apoyo, est, evid = t
        orden[paso] = orden.get(paso, 0) + 1
        rows.append([tid, paso, orden[paso], {1: F1, 2: F2, 3: F3}[paso], fase, nom, acc, ent, cmd, dep, resp, None, apoyo, est, None, evid])
    cols = ["ID", "Paso", "Orden", "Fecha límite", "Fase", "Nombre de la tarea", "Acción / tarea", "Entregable (ruta)", "Comando",
            "Depende de", "Responsable", "Perfil", "Apoyo", "Estado", "¿Hecho?", "Evidencia / observación"]
    ws = sheet(wb, "Matriz", cols, [7, 6, 6, 12, 11, 28, 46, 32, 28, 13, 16, 18, 15, 14, 8, 38], rows)
    last = len(rows) + 1
    for r in range(2, last + 1):
        ws.cell(r, 12, f'=IFERROR(INDEX(Perfiles!$B$2:$B$5,MATCH(K{r},Perfiles!$A$2:$A$5,0)),"")').font = FONT
        ws.cell(r, 15, f'=IF(N{r}="Hecho","Sí","No")').font = FONT
        paso = ws.cell(r, 2).value
        for c in (1, 2, 3):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=PASO_FILL[paso])
    dv = DataValidation(type="list", formula1=f'"{H},{L},{P},En curso,Bloqueado"')
    ws.add_data_validation(dv)
    dv.add(f"N2:N{last}")
    for txt, color in [(H, "C6EFCE"), (L, "DDEBF7"), ("En curso", "FFEB9C"), (P, "F2F2F2"), ("Bloqueado", "FFC7CE")]:
        ws.conditional_formatting.add(f"N2:N{last}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=color)))

    wp = sheet(wb, "Perfiles", ["Integrante", "Perfil", "Responsabilidad", "Tareas", "Hechas", "Listas p/desplegar", "Pendientes", "% hecho"],
               [18, 22, 60, 8, 8, 12, 10, 9], [list(p) for p in PERFILES])
    for r in range(2, 6):
        wp.cell(r, 4, f"=COUNTIF(Matriz!$K$2:$K${last},A{r})")
        for c, st in zip((5, 6, 7), (H, L, P)):
            wp.cell(r, c, f'=COUNTIFS(Matriz!$K$2:$K${last},$A{r},Matriz!$N$2:$N${last},"{st}")')
        wp.cell(r, 8, f"=IF(D{r}>0,E{r}/D{r},0)").number_format = "0%"
        for c in range(4, 9):
            wp.cell(r, c).font, wp.cell(r, c).border = FONT, BRD

    sheet(wb, "Despliegue", ["Paso", "#", "Acción", "Comando", "Responsable", "Insumo del JSON de control", "Verificación"],
          [6, 6, 30, 70, 16, 34, 40], [list(d) for d in DESPLIEGUE])

    wr = wb.create_sheet("Resumen", 0)
    for i, (c, w) in enumerate(zip(["Paso", "Entrega", "Fecha", "Tareas", "Hechas", "Listas p/desplegar", "Pendientes", "% hecho", "% código listo"],
                                   [7, 40, 13, 8, 8, 12, 10, 9, 12]), 1):
        cell = wr.cell(1, i, c)
        cell.font, cell.fill, cell.border = HDR, FILL, BRD
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        wr.column_dimensions[get_column_letter(i)].width = w
    ent = [(1, "Examen + P1: diseño, costos, ingesta/storage", F1), (2, "P2: pipelines batch + streaming, modelado", F2),
           (3, "P3: API, asistente, tableros, demo, sustentación", F3)]
    for r, (paso, nom, fecha) in enumerate(ent, 2):
        wr.cell(r, 1, paso); wr.cell(r, 2, nom); wr.cell(r, 3, fecha)
        wr.cell(r, 4, f"=COUNTIF(Matriz!$B$2:$B${last},A{r})")
        for c, st in zip((5, 6, 7), (H, L, P)):
            wr.cell(r, c, f'=COUNTIFS(Matriz!$B$2:$B${last},$A{r},Matriz!$N$2:$N${last},"{st}")')
        wr.cell(r, 8, f"=IF(D{r}>0,E{r}/D{r},0)")
        wr.cell(r, 9, f"=IF(D{r}>0,(E{r}+F{r})/D{r},0)")
    wr.cell(5, 2, "TOTAL")
    for c in range(4, 8):
        col = get_column_letter(c)
        wr.cell(5, c, f"=SUM({col}2:{col}4)")
    wr.cell(5, 8, "=IF(D5>0,E5/D5,0)"); wr.cell(5, 9, "=IF(D5>0,(E5+F5)/D5,0)")
    for r in range(2, 6):
        for c in range(1, 10):
            wr.cell(r, c).font, wr.cell(r, c).border = (BOLD if r == 5 else FONT), BRD
        wr.cell(r, 8).number_format = wr.cell(r, 9).number_format = "0%"
    notas = ["Leyenda",
             "Paso 1 = Entrega 1 (24-oct) · Paso 2 = Entrega 2 (07-nov) · Paso 3 = repositorio 20-nov y sustentación 21-nov.",
             "Hecho = terminado con evidencia. Listo para desplegar = código probado localmente; falta ejecutarlo en GCP.",
             "Estado se elige en la hoja Matriz (lista desplegable); '¿Hecho?', perfiles y resumen se recalculan solos.",
             "Despliegue: orden de comandos. Parámetros, nombres, tablas, columnas, usuarios y permisos: config/control_plataforma.json (editable con `python ops.py panel`)."]
    for i, t in enumerate(notas):
        wr.cell(8 + i, 1, t).font = BOLD if i == 0 else FONT
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def md_table(cols, rows):
    esc = lambda v: str(v).replace("|", "\\|")
    return "\n".join(["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)] + ["| " + " | ".join(esc(v) for v in r) + " |" for r in rows])


def build_md(out):
    txt = ["# Despliegue en GCP", "",
           "Todos los inputs (parámetros, nombres, tablas, columnas, usuarios iniciales y permisos) están en "
           "`config/control_plataforma.json`. `python ops.py config` los valida y genera las variables de Terraform "
           "(`infra/terraform/plataforma.auto.tfvars.json`), la configuración de ops (`config/generated/plataforma.json`) "
           "y las semillas de la API (`api/seed.json`).", "",
           "## Dónde ejecutar", "",
           "| Paso | Dónde | Por qué |", "|---|---|---|",
           "| subir-datos, local | PC con `data/raw` | Los 3.9 GB de CSV (y 29 GB de imágenes) están ahí |",
           "| todo lo demás | Google Cloud Shell (recomendado) o la PC | Cloud Shell ya trae gcloud, Terraform, Python y Docker y está autenticado |", "",
           "## Pasos", "", md_table(["Paso", "#", "Acción", "Comando", "Responsable", "Insumo del JSON", "Verificación"], DESPLIEGUE), "",
           "## Atajo", "", "```bash", "python ops.py config", "python ops.py todo     # bootstrap, apply, bronze, silver, gold, api, verificar",
           "```", "", "Requiere haber subido antes los datos con `python ops.py subir-datos` desde la PC.", "",
           "## Notas", "",
           "- `plan` antes de `apply`; la CI solo ejecuta `config --check`, pruebas, `fmt` y `validate`, nunca `apply`.",
           "- Dataflow cobra mientras corre: encenderlo solo para la demo.",
           "- El estado de Terraform vive en `gs://<project_id>-tfstate/retail/<env>` (versionado); nunca en Git.",
           "- Contraseñas y tokens no van en config/control_plataforma.json: Terraform los genera en Secret Manager.",
           "- Vertex AI Gemini no es gratuito (~USD 0.0014 por conversación con 3.1 Flash-Lite); `assistant_mode = rules` lo evita.",
           "- `roles/owner` no se puede otorgar por API: quien crea el proyecto ya lo tiene; al resto se les da roles desde la hoja Usuarios."]
    out.write_text("\n".join(txt) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_xlsx(ROOT / "docs/gestion/Matriz_Auditoria_Proyecto.xlsx")
    build_md(ROOT / "docs/DEPLOY.md")
    print(len(TAREAS), "tareas,", len(DESPLIEGUE), "pasos de despliegue,")
